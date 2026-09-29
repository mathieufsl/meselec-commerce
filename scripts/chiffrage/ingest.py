#!/usr/bin/env python3
"""Ingestion de l'historique de chiffrage (réponses AO + brouillons de décomposition).

Usage:
  python3 scripts/chiffrage/ingest.py <dossier "Etude EP"> <dossier "Chiffrage Brouillon"> [--sql <fichier.sql>] [--json <fichier.json>]

- Le PU de la réponse déposée fait foi ; la décomposition (fourniture / MO / marges)
  vient des brouillons, rattachée par numéro de prix.
- Sortie : rapport de couverture + (optionnel) migration SQL idempotente (uuid5 déterministes).
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import uuid
from pathlib import Path

import openpyxl

NS = uuid.UUID("6f0c7a52-3a51-4c2e-9d0e-4d6d6f1a0c11")


# ---------------------------------------------------------------- utilitaires
def norm(s: object) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def num(v: object) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", " ").replace("€", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def txt(v: object) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def sheet_rows(path: Path, sheet: str | int | None = None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet or 0]
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        yield i, list(r)


def sheet_names(path: Path) -> list[str]:
    return openpyxl.load_workbook(path, read_only=True).sheetnames


UNITES = {
    "ml": "ml", "m": "ml", "u": "u", "unite": "u", "unité": "u", "ft": "ft", "fft": "ft",
    "forfait": "ft", "ff": "ft", "m2": "m²", "m²": "m²", "m3": "m³", "m³": "m³", "h": "h",
    "j": "j", "jour": "j", "ens": "ens", "ensemble": "ens", "pm": "pm", "kg": "kg",
    "foyer": "u", "ft/an": "ft/an", "%": "%",
}


def unite_norm(u: str) -> str | None:
    u = txt(u)
    if not u:
        return None
    return UNITES.get(norm(u), UNITES.get(u.lower(), u if len(u) <= 8 else None))


FAMILLES: list[tuple[str, str]] = [
    (r"cout horaire|main d'?oeuvre|taux horaire", "main_oeuvre"),
    (r"nacelle|camion|grue|pelle", "location_materiel"),
    (r"guirlande|motif|illumination|rideau lumineux|traversee de (rue|chaussee)|plafond lumineux|chapiteau|decorations", "illuminations"),
    (r"astreinte", "astreinte"),
    (r"tournee", "tournee"),
    (r"maintenance|entretien|exploitation|curatif|preventif|controle", "maintenance"),
    (r"gmao|base de donnees|geolocalis|georef|plan de recolement|plans d'execution|rapport", "etudes_donnees"),
    (r"location", "location_materiel"),
    (r"tranchee|fouille|terrassement|demolition|remblai|enrobe|carottage|pavage|beton", "genie_civil"),
    (r"fourreau|gaine|cable|cuivre|aluminium|conducteur|grillage", "reseaux"),
    (r"candelabre|mat |mat$|crosse|console|potence|fondation|potelet", "supports"),
    (r"luminaire|projecteur|lanterne|lampe|led |relamping|point lumineux", "luminaires"),
    (r"armoire|coffret|tableau|disjoncteur|horloge|contacteur|tgbt|regard", "armoires_coffrets"),
    (r"installation de chantier|signalisation|balise|barriere|forfait d'intervention|preparation|autorisation|huissier|etude|enquete|sondage|marquage", "prestations_chantier"),
]


def famille_de(designation: str, section: str = "") -> str | None:
    n = norm(f"{designation} {section}")
    for pat, fam in FAMILLES:
        if re.search(pat, n):
            return fam
    return None


DIM_RE = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*m?\s*[x×*]\s*(\d+(?:[.,]\d+)?)\s*m?\b", re.I
)


def dimensions_de(designation: str) -> dict:
    m = DIM_RE.search(txt(designation))
    if not m:
        return {}
    a, b = (float(g.replace(",", ".")) for g in m.groups())
    if a > 40 or b > 40:  # sections cable / mm : pas des dimensions de motif
        return {}
    return {"largeur_m": a, "hauteur_m": b}


# ------------------------------------------------- parseur générique de tableaux
ROLE_PATTERNS = [
    ("total", r"^(p\.?t\.?\b|prix total|total|montant)"),
    ("pu", r"^(p\.?\s?u\.?\b|p unitaire|prix\s*u\b|prix unit|tarif)"),
    ("qte", r"^(q|qt|qte|quantite)\b"),
    ("unite", r"^(u|unite)$"),
    ("numero", r"^(n|no|n article|n prix|n de prix|ref|art|ref\.? bpu)\b"),
    ("designation", r"(designation|libelle|type de motif|type de decor)"),
]
FILLER = re.compile(r"^(l.?unite|par an|le forfait|selon cctp|forfait|unite)\s*:?\s*$")
NUMERIC_CODE = re.compile(r"^\d+(\.\d+)*$")


def detect_header(row: list) -> list[tuple[int, str]] | None:
    roles: list[tuple[int, str]] = []
    for c, v in enumerate(row):
        t = norm(v)
        if not t or len(t) > 40 or "ttc" in t:
            continue
        for role, pat in ROLE_PATTERNS:
            if re.search(pat, t):
                roles.append((c, role))
                break
    names = {r for _, r in roles}
    if "pu" in names and ("designation" in names or "unite" in names or "numero" in names):
        return roles
    return None


def split_blocks(roles: list[tuple[int, str]]) -> list[dict[str, int]]:
    blocks: list[dict[str, int]] = []
    cur: dict[str, int] = {}
    for c, role in roles:
        if role in cur:
            blocks.append(cur)
            cur = {}
        cur[role] = c
    if cur:
        blocks.append(cur)
    return [b for b in blocks if "pu" in b]


def depth_of(numero: str) -> int:
    n = norm(numero)
    if n.startswith("chap"):
        return 0
    if n.startswith("rubrique"):
        return 1
    if NUMERIC_CODE.match(numero):
        return numero.count(".") + 1
    return 1


def parse_generic(path: Path, sheet: str | int | None, *, fichier: str) -> list[dict]:
    """Détecte les en-têtes (plusieurs possibles, blocs côte à côte) et lit les lignes chiffrées."""
    out: list[dict] = []
    blocks: list[dict[str, int]] = []
    stack: dict[int, str] = {}
    pending: dict[int, tuple[str, str]] = {}  # bloc -> (numero, titre) en attente d'un prix
    sheet_label = sheet if isinstance(sheet, str) else str(sheet or 0)
    for i, row in sheet_rows(path, sheet):
        hdr = detect_header(row)
        if hdr:
            blocks = split_blocks(hdr)
            pending = {}
            continue
        if not blocks:
            continue
        for bi, b in enumerate(blocks):
            def cell(role: str, _b=b, _row=row):
                c = _b.get(role)
                return _row[c] if c is not None and c < len(_row) else None

            pu, qte, total = num(cell("pu")), num(cell("qte")), num(cell("total"))
            des = txt(cell("designation"))
            if not des and "designation" not in b:
                anchor = b.get("unite", b["pu"])
                for c in range(anchor - 1, -1, -1):
                    if c < len(row) and isinstance(row[c], str) and txt(row[c]):
                        des = txt(row[c])
                        break
            numero = txt(cell("numero"))
            if (not numero or len(numero) == 1) and b.get("numero", 0) != 0 and NUMERIC_CODE.match(txt(row[0])):
                numero = txt(row[0])  # numéro non étiqueté en colonne A (ex : BPU Jouars)
            unite = unite_norm(cell("unite"))

            if pu is not None and pu > 0 and (des or numero):
                pend_num, pend_titre = pending.get(bi, ("", ""))
                used_pending = False
                if not numero and pend_num:
                    numero, used_pending = pend_num, True
                if FILLER.match(norm(des)):
                    des = ""
                full = des
                if pend_titre and (used_pending or not numero) and norm(pend_titre) not in norm(des):
                    full = f"{pend_titre} — {des}" if des else pend_titre
                elif used_pending and not des:
                    full = pend_titre
                skip = depth_of(numero) if used_pending else None
                section = " > ".join(t for d, t in sorted(stack.items()) if d != skip)
                out.append({
                    "numero": numero or None, "section": section or None,
                    "designation": full, "unite": unite,
                    "quantite": qte if qte and qte > 0 else None, "pu_ht": pu,
                    "montant_ht": total if total and total > 0 else None,
                    "fichier": fichier, "feuille": sheet_label, "ligne": i,
                })
                pending.pop(bi, None)
            elif des and numero and pu is None and unite is None:
                d = depth_of(numero)
                stack = {k: v for k, v in stack.items() if k < d}
                stack[d] = des
                pending[bi] = (numero, des)
            elif des and pu is None and unite is None and not numero and not any(
                num(row[c]) is not None for c in b.values() if c < len(row)
            ):
                if re.match(r"^\d+\s*-\s*", des) or (des.isupper() and len(des) > 12):
                    stack = {0: des}
    return out


# ------------------------------------------------------- décomposition (brouillons)
def parse_devis(path: Path, sheet: str, *, fichier: str) -> tuple[list[dict], float | None]:
    """Gabarit « devis » MESELEC (Jouars, Champagne…) : PU + fourniture / ST / heures MO."""
    rows = list(sheet_rows(path, sheet))
    hdr_i = next(i for i, r in rows if norm(r[0]) == "n" and "designation travaux" in norm(r[1]))
    hdr = next(r for i, r in rows if i == hdr_i)
    col = {norm(v): c for c, v in enumerate(hdr) if v is not None}

    def find(*keys: str) -> int | None:
        for k in keys:
            for name, c in col.items():
                if name.startswith(k):
                    return c
        return None

    c_pu, c_qt, c_u = find("p unitaire"), find("qt"), find("u")
    c_fo_n, c_st_n, c_st_p = find("fo"), find("st"), find("prix u st")
    c_fourn, c_stt, c_dir = find("fourniture en euros"), find("st en euros"), find("heure direction")
    c_h = [c for name, c in col.items() if name.startswith("h total") or name.startswith("h totalde")]
    thm = None
    out: list[dict] = []
    pend: tuple[str, str] | None = None
    stack: dict[int, str] = {}
    for i, r in rows:
        if i <= hdr_i:
            continue
        numero, des = txt(r[0]), txt(r[1])
        pu = num(r[c_pu]) if c_pu is not None else None
        unite = unite_norm(r[c_u]) if c_u is not None and isinstance(r[c_u], str) else None
        if pu and pu > 0.005 and unite:
            titre = ""
            if not numero and pend:
                numero, titre = pend
            fourn = num(r[c_fourn]) if c_fourn is not None else None
            st = num(r[c_stt]) if c_stt is not None else None
            heures = round(sum((num(r[c]) or 0) for c in c_h if c < len(r) and (num(r[c]) or 0) > 1e-6), 4)
            full = f"{titre} — {des}" if titre and norm(titre) not in norm(des) and not FILLER.match(norm(des)) else (des or titre)
            out.append({
                "numero": numero or None, "designation": full, "unite": unite,
                "quantite": num(r[c_qt]) if c_qt is not None else None,
                "pu_brouillon_ht": pu,
                "fourniture_ht": fourn if fourn and fourn > 1e-6 else None,
                "sous_traitance_ht": st if st and st > 1e-6 else None,
                "fournisseur": txt(r[c_fo_n]) or None if c_fo_n is not None and isinstance(r[c_fo_n], str) else None,
                "sous_traitant": txt(r[c_st_n]) or None if c_st_n is not None and isinstance(r[c_st_n], str) and num(r[c_st_n]) is None else None,
                "heures_mo": heures or None,
                "detail": f"H direction BU : {round(num(r[c_dir]) or 0, 3)}" if c_dir is not None and (num(r[c_dir]) or 0) > 1e-3 else None,
                "fichier": fichier, "feuille": sheet, "ligne": i,
            })
            pend = None
        elif numero and des and not unite:
            pend = (numero, des)
    # THM (feuille de vente, ligne « Direction BU »)
    try:
        for i, r in sheet_rows(path, "FV"):
            if txt(r[1]).lower().startswith("direction bu"):
                thm = num(r[3])
                break
    except KeyError:
        pass
    if thm:
        for o in out:
            if o.get("heures_mo"):
                o["mo_ht"] = round(o["heures_mo"] * thm, 4)
    return out, thm


def parse_goussainville_detail(path: Path) -> list[dict]:
    out = []
    for i, r in sheet_rows(path, "Detail_chiffrage"):
        if i == 1 or not r[0]:
            continue
        pu = num(r[6])
        out.append({
            "numero": txt(r[0]), "designation": txt(r[1]), "unite": unite_norm(r[2]),
            "fourniture_ht": num(r[3]) or None, "heures_mo": num(r[4]) or None,
            "mo_ht": num(r[5]) or None, "pu_brouillon_ht": pu,
            "aleas_ht": num(r[8]), "frais_generaux_ht": num(r[9]), "marge_nette_ht": num(r[10]),
            "fichier": path.name, "feuille": "Detail_chiffrage", "ligne": i,
        })
    return out


def attach_decomposition(postes: list[dict], deco: list[dict], origine: str) -> tuple[int, list[str]]:
    """Rattache la décomposition aux postes de la réponse par numéro (contrôle sur la désignation)."""
    by_num: dict[str, list[dict]] = {}
    for p in postes:
        if p.get("numero"):
            by_num.setdefault(p["numero"], []).append(p)
    ok, warns = 0, []
    for d in deco:
        cands = [p for p in by_num.get(d["numero"] or "", []) if "decomposition_origine" not in p]
        if not cands:
            warns.append(f"{origine} n°{d['numero']} sans poste correspondant dans la réponse")
            continue
        p = cands[0]
        a, b = norm(p["designation"])[:24], norm(d["designation"])[:24]
        if a and b and a[:12] != b[:12] and a not in norm(d["designation"]) and b not in norm(p["designation"]):
            warns.append(f"{origine} n°{d['numero']} : désignation différente ({a!r} / {b!r})")
        for k in ("fourniture_ht", "heures_mo", "mo_ht", "sous_traitance_ht", "fournisseur", "sous_traitant",
                  "aleas_ht", "frais_generaux_ht", "marge_nette_ht", "pu_brouillon_ht"):
            if d.get(k) is not None:
                p[k] = d[k]
        if d.get("detail"):
            p["detail"] = d["detail"]
        p["decomposition_origine"] = origine
        ok += 1
    return ok, warns


def merge_quantites(base: list[dict], dqe: list[dict]) -> int:
    """Reporte quantités / montants d'un DQE sur les lignes du BPU (clé : numéro, sinon désignation + PU)."""
    used: set[int] = set()
    n = 0
    for q in dqe:
        for idx, b in enumerate(base):
            if idx in used or b.get("quantite"):
                continue
            same_num = q.get("numero") and q["numero"] == b.get("numero")
            same_txt = norm(q["designation"])[:50] == norm(b["designation"])[:50]
            if (same_num or same_txt) and abs(q["pu_ht"] - b["pu_ht"]) <= max(0.02, 0.005 * b["pu_ht"]):
                b["quantite"], b["montant_ht"] = q.get("quantite"), q.get("montant_ht")
                used.add(idx)
                n += 1
                break
        else:
            base.append({**q, "detail": "ligne du DQE absente du BPU"})
    return n


# ---------------------------------------------------------------- profils atypiques
def parse_arc(root: Path, ch: Path) -> list[dict]:
    rep = root / "Arc de Triomphe/03_Réponse/Offre"
    out: list[dict] = []
    f = rep / "26-641-100 - DPGF_MESELEC.xlsx"
    for i, r in sheet_rows(f, "DPGF"):
        pu = num(r[5])
        if pu and txt(r[0]) and not norm(r[0]).startswith(("total", "centre", "exploitation", "pour")):
            out.append({"numero": None, "section": "DPGF forfaits annuels", "designation": txt(r[0]),
                        "unite": "ft/an", "quantite": 1.0, "pu_ht": pu, "montant_ht": pu,
                        "fichier": f.name, "feuille": "DPGF", "ligne": i})
    f = rep / "26-641-100 - BPU-DQE_MESELEC.xlsx"
    for i, r in sheet_rows(f, "BPU"):
        if txt(r[1]) and num(r[4]) and i < 14:
            out.append({"numero": None, "section": "Coûts horaires", "designation": f"Coût horaire {txt(r[1])} (lun-ven 7h-18h)",
                        "unite": "h", "pu_ht": num(r[4]), "fichier": f.name, "feuille": "BPU", "ligne": i})
        elif txt(r[3]).startswith("Z") and num(r[4]):
            out.append({"numero": txt(r[3]), "section": "Taux de marge sur fournitures", "designation": txt(r[0]),
                        "unite": "coef", "pu_ht": num(r[4]), "fichier": f.name, "feuille": "BPU", "ligne": i})
    for i, r in sheet_rows(f, "DQE"):
        qte, taux, tot = num(r[2]), num(r[3]), num(r[4])
        if txt(r[1]) and qte and taux:
            out.append({"numero": None, "section": "DQE annuel", "designation": txt(r[1]), "unite": "h",
                        "quantite": qte, "pu_ht": taux, "montant_ht": tot,
                        "fichier": f.name, "feuille": "DQE", "ligne": i})
    # Hypothèses du brouillon « direction » (41 400 € ≠ offre déposée 36 010 €) : conservées en note
    hyp = ch / "OneDrive_11_29-09-2026" / "DPGF_Arc_de_Triomphe_chiffrage_detaille_direction.xlsx"
    if hyp.exists():
        deb = {norm(r[0]): r for i, r in sheet_rows(hyp, "Déboursé MO") if i > 2 and r[0]}
        for o in out:
            if o["section"] != "DPGF forfaits annuels":
                continue
            d = deb.get(norm(o["designation"]))
            if d:
                o["heures_mo"] = num(d[3]) or None
                o["detail"] = (f"Brouillon direction : {txt(d[1])}, {txt(d[2])}, {num(d[3]) or 0:g} h × {num(d[4]) or 0:g} €/h "
                               f"+ {num(d[5]) or 0:g} € autres = {num(d[6]) or 0:g} € (offre déposée : {o['pu_ht']:g} €)")
                o["pu_brouillon_ht"] = num(d[6])
                o["decomposition_origine"] = "arc_deboursé_mo"
    return out


def parse_moissy(f: Path) -> list[dict]:
    out, lieu = [], ""
    for i, r in sheet_rows(f, 0):
        if i < 7 or not txt(r[1]) or norm(r[0]).startswith(("sous total", "total", "tva")):
            continue
        lieu = txt(r[0]) or lieu
        des = txt(r[1])
        forfait, loc = num(r[2]), num(r[3])
        if forfait:
            out.append({"numero": None, "section": lieu, "designation": f"Pose, entretien, dépose — {des}", "unite": "ft/an",
                        "quantite": 1.0, "pu_ht": forfait, "montant_ht": forfait, "fichier": f.name, "feuille": "DPGF", "ligne": i})
        if loc:
            out.append({"numero": None, "section": lieu, "designation": f"Location motif — {des}", "unite": "ft/an",
                        "quantite": 1.0, "pu_ht": loc, "montant_ht": loc, "fichier": f.name, "feuille": "DPGF", "ligne": i})
    return out


def parse_epinay(f: Path) -> list[dict]:
    out = []
    for i, r in sheet_rows(f, 0):
        pu, qte, tot = num(r[3]), num(r[5]), num(r[6])
        t = txt(r[0])
        if pu and t and qte is not None and not norm(t).startswith(("taux", "type")):
            dims = [txt(r[1]), txt(r[2])]
            dim_txt = " x ".join(d for d in dims if d and d not in ("/", ""))
            out.append({"numero": None, "section": "Décors lumineux de Noël 2026-2029",
                        "designation": f"{t} {dim_txt}".strip(), "unite": "u", "quantite": qte, "pu_ht": pu,
                        "montant_ht": tot, "fichier": f.name, "feuille": "BPU", "ligne": i})
    return out


def parse_aubervilliers(f: Path) -> list[dict]:
    base = parse_generic(f, "BPU", fichier=f.name)
    dqe = parse_generic(f, "DQE", fichier=f.name)
    fourn = {}
    for b in base + dqe:
        m = re.match(r"^F(\d+)$", b.get("numero") or "")
        if m:
            fourn[m.group(1)] = b["designation"]
    for b in base + dqe:
        m = re.match(r"^PDM(\d+)$", b.get("numero") or "")
        if m:
            ref = fourn.get(m.group(1), "")
            b["designation"] = f"Pose, dépose et maintenance — {ref}" if ref else f"Pose, dépose et maintenance {b['numero']}"
    merge_quantites(base, dqe)
    return base


def jouars_hypotheses(f: Path) -> dict[str, str]:
    """Commentaires d'hypothèses du DPGF « HRB » (colonne commentaires), par numéro de prix."""
    out, numero = {}, ""
    for i, r in sheet_rows(f, "DPGF"):
        if i < 3:
            continue
        numero = txt(r[0]) or numero
        if len(r) > 7 and txt(r[7]) and numero:
            out[numero] = txt(r[7])
    return out


# ------------------------------------------------------------------- assemblage
SOURCES = [
    dict(code="osny", nom="Osny — éclairage, illuminations et enfouissement", client="Ville d'Osny", activite="ep", ao_ref="AO-OSNY"),
    dict(code="jouars", nom="Jouars-Pontchartrain — bail EP et SLT", client="Ville de Jouars-Pontchartrain", activite="ep", ao_ref="AO-JOUARS"),
    dict(code="champagne", nom="Champagne-sur-Oise — éclairage 2026", client="Ville de Champagne-sur-Oise", activite="ep", ao_ref="AO-CHAMP"),
    dict(code="arc-triomphe", nom="Arc de Triomphe — exploitation / maintenance", client="Centre des monuments nationaux", activite="ep", ao_ref="26-641"),
    dict(code="goussainville-l2", nom="Goussainville — Lot 2", client="Ville de Goussainville", activite="ep", ao_ref="260713"),
    dict(code="aubervilliers", nom="Aubervilliers — illuminations", client="Ville d'Aubervilliers", activite="illuminations", ao_ref=None),
    dict(code="epinay", nom="Épinay-sur-Seine — décors de Noël 2026-2029", client="Ville d'Épinay-sur-Seine", activite="illuminations", ao_ref=None),
    dict(code="moissy", nom="Moissy-Cramayel — illuminations 2026S18", client="Commune de Moissy-Cramayel", activite="illuminations", ao_ref=None),
    dict(code="paris-11-12", nom="Paris 11e/12e — illuminations de Noël", client="Ville de Paris", activite="illuminations", ao_ref=None),
]


def build(ep: Path, brouillon: Path) -> tuple[list[dict], dict[str, list[dict]], list[str]]:
    warns: list[str] = []
    P: dict[str, list[dict]] = {}
    ch = brouillon / "Chiffrage Brouillon" if (brouillon / "Chiffrage Brouillon").exists() else brouillon
    root = ep / "Etude EP" if (ep / "Etude EP").exists() else ep

    # Osny : BPU + quantités du DQE
    o = root / "Osny/Réponse/Offre"
    bpu = parse_generic(next(o.glob("3_*BPU*.xlsx")), 0, fichier=next(o.glob("3_*BPU*.xlsx")).name)
    dqe = parse_generic(next(o.glob("4_*DQE*.xlsx")), 0, fichier=next(o.glob("4_*DQE*.xlsx")).name)
    merge_quantites(bpu, dqe)
    P["osny"] = bpu

    # Jouars : DPGF (forfaits) + BPU (séries) de la réponse ; décomposition = devis « HRB »
    j = root / "Jouars Pontchartrain/04_Réponse_28_05/Offre"
    dpgf = parse_generic(j / "MESELEC_DPGF ECLAIRAGE.xlsx", "DPGF", fichier="MESELEC_DPGF ECLAIRAGE.xlsx")
    for p in dpgf:
        p["section"] = "DPGF forfaits annuels"
    bpu_j = parse_generic(j / "MESELEC_BPU v4_precision du 29042026.xlsx", "Feuil1", fichier="MESELEC_BPU v4_precision du 29042026.xlsx")
    hyp = jouars_hypotheses(next((ch / "OneDrive_10_29-09-2026").glob("DPGF ECLAIRAGE HRB.xlsx")))
    for p in dpgf:
        if p.get("numero") in hyp:
            p["detail"] = f"Hypothèse : {hyp[p['numero']]}"
    for p in dpgf:
        if p.get("numero"):
            p["numero"] = f"DPGF-{p['numero']}"
    devis, _thm = parse_devis(next((ch / "OneDrive_10_29-09-2026").glob("BPU_Jouars*HRB.xlsx")), "devis", fichier="BPU_Jouars pontchartrains_V2 HRB.xlsx")
    n, w = attach_decomposition(bpu_j, devis, "devis")
    warns += [f"jouars: {x}" for x in w[:8]] + ([f"jouars: … {len(w)} avertissements de rattachement"] if len(w) > 8 else [])
    P["jouars"] = dpgf + bpu_j

    # Champagne : pas de réponse xlsx → devis (PU du devis = devis estimatif déposé) + DPGF forfaits + quantités
    devis_c, _ = parse_devis(ch / "OneDrive_12_29-09-2026/BPU Champagne sur oise.xlsx", "devis", fichier="BPU Champagne sur oise.xlsx")
    champ = []
    for d in devis_c:
        champ.append({**d, "section": None, "pu_ht": d["pu_brouillon_ht"], "decomposition_origine": "devis",
                      "detail": "Réponse xlsx absente : PU issu du devis (identique au devis estimatif non contractuel)"})
    est = parse_generic(ch / "OneDrive_12_29-09-2026/Devis Estimatif non contractuel 2026.xlsx", 0, fichier="Devis Estimatif non contractuel 2026.xlsx")
    for e in est:
        for c in champ:
            if e["numero"] and c["numero"] == e["numero"] and not c.get("quantite_estimee"):
                c["quantite"], c["montant_ht"], c["quantite_estimee"] = e.get("quantite"), e.get("montant_ht"), True
                break
    for c in champ:
        c.pop("quantite_estimee", None)
        if c.get("quantite") == 1.0 and not c.get("montant_ht"):
            c["quantite"] = None
    forfaits = parse_generic(ch / "OneDrive_12_29-09-2026/DPGF ECLAIRAGE 2026.xlsx", 0, fichier="DPGF ECLAIRAGE 2026.xlsx")
    for f_ in forfaits:
        f_["section"] = "DPGF forfaits annuels"
        f_["numero"] = f"DPGF-{f_['numero']}" if f_.get("numero") else None
    P["champagne"] = forfaits + champ

    P["arc-triomphe"] = parse_arc(root, ch)

    # Goussainville : BPU réponse + Detail_chiffrage
    g = next((root / "Goussainville/03_Réponse/Offre").glob("*Lot2*MESELEC.xlsx"))
    gb = parse_generic(g, 0, fichier=g.name)
    n, w = attach_decomposition(gb, parse_goussainville_detail(ch / "BPU_Lot2_Goussainville_V4B.xlsx"), "goussainville_detail")
    warns += [f"goussainville: {x}" for x in w[:8]]
    P["goussainville-l2"] = gb

    P["aubervilliers"] = parse_aubervilliers(next((root / "Aubervilliers/Réponse Aubervilliers/OFFRE").glob("*.xlsx")))
    P["epinay"] = parse_epinay(next((root / "Epinay/REPONSE_070926/OFFRE").glob("*.xlsx")))
    P["moissy"] = parse_moissy(next((root / "Moissy-Cramayel/REPONSE/OFFRE").glob("*.xlsx")))
    paris = root / "Paris XI/Réponse EP Paris 11e et 12e/OFFRE/Paris_2600996_Chiffrage_MESELEC.xlsx"
    P["paris-11-12"] = parse_generic(paris, "Facture-type_a_deposer", fichier=paris.name)

    for code in P:
        P[code] = [r for r in P[code] if not re.match(r"^(total|sous.?total)\b", norm(r["designation"])) and "pour information" not in norm(r["designation"])]
    for code, rows in P.items():
        for k, r in enumerate(rows):
            r["ordre"] = k
            r["designation"] = txt(r["designation"])[:600]
            r["famille"] = famille_de(r["designation"], r.get("section") or "")
            r["dimensions"] = dimensions_de(r["designation"])
            q, pu, m = r.get("quantite"), r.get("pu_ht"), r.get("montant_ht")
            if q and pu and m and abs(q * pu - m) > max(0.05, 0.01 * m):
                warns.append(f"{code} l.{r.get('ligne')} : quantité × PU ({q*pu:.2f}) ≠ montant ({m:.2f})")
    return SOURCES, P, warns


# ------------------------------------------------------------------------ sorties
COLS = ["numero", "section", "designation", "unite", "quantite", "pu_ht", "montant_ht", "famille", "dimensions",
        "fourniture_ht", "heures_mo", "mo_ht", "sous_traitance_ht", "fournisseur", "sous_traitant", "aleas_ht",
        "frais_generaux_ht", "marge_nette_ht", "pu_brouillon_ht", "decomposition_origine", "detail", "fichier", "feuille",
        "ligne", "ordre"]


def q(v: object) -> str:
    if v is None:
        return "null"
    if isinstance(v, (dict, list)):
        return "'" + json.dumps(v, ensure_ascii=False).replace("'", "''") + "'::jsonb"
    if isinstance(v, (int, float)):
        return repr(round(float(v), 6)) if isinstance(v, float) else str(v)
    return "'" + str(v).replace("'", "''") + "'"


def to_sql(sources: list[dict], P: dict[str, list[dict]]) -> str:
    L = ["-- Généré par scripts/chiffrage/ingest.py — historique de chiffrage (offres déposées + décomposition).",
         "-- Idempotent : les postes des sources listées sont remplacés à chaque exécution.", ""]
    for s in sources:
        sid = uuid.uuid5(NS, f"source:{s['code']}")
        fichiers = sorted({r["fichier"] for r in P[s["code"]] if r.get("fichier")})
        L.append(
            "insert into public.chiffrage_sources (id, code, nom, client, activite, annee, ao_id, fichier_reponse) values "
            f"({q(str(sid))}, {q(s['code'])}, {q(s['nom'])}, {q(s['client'])}, {q(s['activite'])}, 2026, "
            + (f"(select id from public.appels_offres where reference = {q(s['ao_ref'])} limit 1)" if s.get("ao_ref") else "null")
            + f", {q(' ; '.join(fichiers))}) on conflict (code) do update set nom = excluded.nom, client = excluded.client, "
            "activite = excluded.activite, fichier_reponse = excluded.fichier_reponse;"
        )
    L.append("")
    ids = ", ".join(q(str(uuid.uuid5(NS, f"source:{s['code']}"))) for s in sources)
    L.append(f"delete from public.chiffrage_postes where source_id in (select id from public.chiffrage_sources where code in ({', '.join(q(s['code']) for s in sources)}));")
    for s in sources:
        rows = P[s["code"]]
        if not rows:
            continue
        for start in range(0, len(rows), 200):
            chunk = rows[start:start + 200]
            vals = ",\n".join(
                "((select id from public.chiffrage_sources where code = " + q(s["code"]) + "), "
                + ", ".join(q(r.get(c)) for c in COLS) + ")"
                for r in chunk
            )
            L.append(f"insert into public.chiffrage_postes (source_id, {', '.join(COLS)}) values\n{vals};")
    return "\n".join(L) + "\n"


def report(sources: list[dict], P: dict[str, list[dict]], warns: list[str]) -> None:
    print(f"{'source':18} {'postes':>6} {'qté':>5} {'décomp.':>8} {'famille?':>9} {'unité?':>7}")
    for s in sources:
        r = P[s["code"]]
        n = len(r) or 1
        print(f"{s['code']:18} {len(r):6} {sum(1 for x in r if x.get('quantite')):5} "
              f"{sum(1 for x in r if x.get('decomposition_origine')):8} {sum(1 for x in r if not x.get('famille')):9} "
              f"{sum(1 for x in r if not x.get('unite')):7}")
    print(f"\nTotal postes : {sum(len(v) for v in P.values())}")
    for w in warns[:40]:
        print("  ⚠", w)
    if len(warns) > 40:
        print(f"  … {len(warns) - 40} avertissements supplémentaires")


def main() -> None:
    args = sys.argv[1:]
    if len(args) < 2:
        sys.exit(__doc__)
    sql_out = args[args.index("--sql") + 1] if "--sql" in args else None
    json_out = args[args.index("--json") + 1] if "--json" in args else None
    sources, P, warns = build(Path(args[0]), Path(args[1]))
    report(sources, P, warns)
    if json_out:
        Path(json_out).write_text(json.dumps(P, ensure_ascii=False, indent=1))
    if sql_out:
        Path(sql_out).write_text(to_sql(sources, P))
        print(f"\nSQL écrit : {sql_out}")


if __name__ == "__main__":
    main()
